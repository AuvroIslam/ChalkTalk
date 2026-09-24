"""Knowledge beyond the pixels: what the user is looking at, and tools to read more.

When the hotkey is pressed we work out the source (a YouTube video, a PDF or
slide deck on disk, a web page) and start loading it in the background while
the user types. The model gets a short digest up front (often enough to answer
straight away) plus typed tools to pull in the full transcript, other pages,
the whole web page, or a web search when the screen alone isn't enough.
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import tempfile
import threading
import urllib.parse
from concurrent.futures import Future, ThreadPoolExecutor
from pathlib import Path
from typing import Callable

import httpx
import openai
from pydantic import BaseModel, Field, ValidationError

import config
from window import Foreground, browser_url

MAX_TOOL_CHARS = 24000
DOC_EXTS = (".pdf", ".pptx", ".docx", ".txt", ".md")
_pool = ThreadPoolExecutor(max_workers=3, thread_name_prefix="context")
_WORD = re.compile(r"[\w']{3,}", re.UNICODE)


def _words(text: str) -> set[str]:
    return {w.lower()[:6] for w in _WORD.findall(text)}  # crude stemming: "gradients" ~ "gradient"


def _clip(text: str, n: int = MAX_TOOL_CHARS) -> str:
    return text if len(text) <= n else text[:n] + f"\n…[cut: {len(text) - n} more characters]"


def _mmss(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 3600}:{sec % 3600 // 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60}:{sec % 60:02d}"


def _seconds(s: str) -> float:
    parts = [int(p) for p in s.split(":")]
    out = 0
    for p in parts:
        out = out * 60 + p
    return float(out)


_emb: dict[str, "np.ndarray"] = {}
_emb_off = False


def embed(texts: list[str]):
    """Embeddings for texts (cached), or None when unavailable -> keyword search only."""
    global _emb_off
    import hashlib

    import numpy as np

    if _emb_off or not config.EMBED_MODEL or not texts:
        return None
    keys = [hashlib.sha1(t.encode("utf-8", "replace")).hexdigest() for t in texts]
    todo = [(k, t) for k, t in zip(keys, texts) if k not in _emb]
    try:
        from clients import make_client

        for i in range(0, len(todo), 96):
            batch = todo[i:i + 96]
            r = make_client().embeddings.create(model=config.EMBED_MODEL,
                                                input=[(t[:6000] or " ") for _, t in batch])
            for (k, _), d in zip(batch, r.data):
                v = np.asarray(d.embedding, dtype=np.float32)
                _emb[k] = v / (np.linalg.norm(v) + 1e-9)
    except Exception as e:
        print(f"  embeddings unavailable ({type(e).__name__}: {str(e)[:80]}); using keyword search")
        _emb_off = True
        return None
    return np.stack([_emb[k] for k in keys])


def _chunk_text(label: str, text: str) -> str:
    return f"{label}\n{text}" if label else text


def index(chunks: list[tuple[str, str]]) -> None:
    """Embed chunks ahead of time (in the background) so searching them later is instant."""
    embed([_chunk_text(lab, t) for lab, t in chunks])


def _best_windows(chunks: list[tuple[str, str]], query: str, k: int = 5) -> list[tuple[str, str]]:
    """Rank (label, text) chunks for a query: by meaning (embeddings) plus the query
    words they contain (rare words count more). Meaning matters: an article may
    call 'variants' "Extensions" or "Momentum"."""
    import math

    if not chunks:
        return []
    q = _words(query)
    sets = [_words(t) for _, t in chunks]
    n = len(sets)
    idf = {w: math.log((n + 1) / (1 + sum(w in s for s in sets))) + 0.1 for w in q}
    kw = [sum(idf[w] for w in q & s) for s in sets]
    top_kw = max(kw) or 1.0
    vecs = embed([_chunk_text(lab, t) for lab, t in chunks] + [query])
    if vecs is not None:
        sims = vecs[:-1] @ vecs[-1]
        scores = [float(s) + 0.25 * (w / top_kw) for s, w in zip(sims, kw)]
    else:
        scores = kw
        if max(kw) <= 0:
            return []
    order = sorted(range(n), key=lambda i: -scores[i])
    return [chunks[i] for i in order[:k]]


# ---- typed tool arguments (strict: the model must fill every field) ----------

class GetVideoTranscript(BaseModel):
    """Read the video's captions between two times. Use for 'what did they say before/after this'."""
    start_sec: float = Field(description="start time in seconds; 0 for the beginning")
    end_sec: float | None = Field(description="end time in seconds; null = end of video")
    reason: str = Field(description="one short line: what context is missing and why you need it")


class SearchVideoTranscript(BaseModel):
    """Find where in the video a topic, term or symbol is talked about. Returns timestamped excerpts."""
    query: str
    reason: str = Field(description="one short line: what context is missing and why you need it")


class ReadDocumentPages(BaseModel):
    """Read pages (or slides) of the open document. Max 20 pages per call."""
    start_page: int = Field(description="1-based")
    end_page: int = Field(description="1-based, inclusive")
    reason: str = Field(description="one short line: what context is missing and why you need it")


class SearchDocument(BaseModel):
    """Find the pages of the open document that mention a term, symbol or topic (e.g. where it's defined)."""
    query: str
    reason: str = Field(description="one short line: what context is missing and why you need it")


class ReadWebpage(BaseModel):
    """Read a web page (null url = the page on screen). Long pages: give `find` to get the sections about it."""
    url: str | None
    find: str | None = Field(description="what to look for on the page, e.g. 'variants and extensions'; null = the start")
    reason: str = Field(description="one short line: what context is missing and why you need it")


class WebSearch(BaseModel):
    """Search the web. Use for facts you are not sure of, recent or niche topics, or to check a claim."""
    query: str
    reason: str = Field(description="one short line: what context is missing and why you need it")


# ---- web (Firecrawl) -----------------------------------------------------------

def _firecrawl(path: str, body: dict) -> dict:
    if not config.FIRECRAWL_KEY:
        raise RuntimeError("web access needs FIRECRAWL_API_KEY in .env")
    r = httpx.post(f"https://api.firecrawl.dev/v2/{path}", json=body, timeout=45,
                   headers={"Authorization": f"Bearer {config.FIRECRAWL_KEY}"})
    r.raise_for_status()
    return r.json().get("data") or {}


_pages: dict[str, tuple[str, str]] = {}  # url -> (title, markdown); pages are scraped once


def _sections(md: str) -> list[str]:
    """Split markdown at headings; very long sections are cut into ~6000-char parts."""
    out = []
    for s in re.split(r"\n(?=#{1,4} )", md):
        head = s.splitlines()[0] if s.startswith("#") else ""
        for i in range(0, len(s), 6000):
            out.append(s[i:i + 6000] if i == 0 else f"{head} (cont.)\n{s[i:i + 6000]}")
    return [s for s in out if s.strip()]


def _scrape(url: str) -> tuple[str, str]:
    if url not in _pages:
        data = _firecrawl("scrape", {"url": url, "formats": ["markdown"], "onlyMainContent": True})
        _pages[url] = ((data.get("metadata") or {}).get("title", ""), data.get("markdown", ""))
    return _pages[url]


def read_url(url: str, find: str | None = None) -> str:
    """The start of the page, or (with `find`) the sections most about it. Long
    pages don't fit in one read: the part the user asks about may be 50k characters in."""
    title, md = _scrape(url)
    head = f"# {title}\n{url}\n"
    sections = _sections(md)
    toc = [s.splitlines()[0].lstrip("# ").strip() for s in sections
           if s.startswith("#") and not s.splitlines()[0].endswith("(cont.)")]
    toc_text = "Sections: " + " | ".join(toc[:60]) + "\n\n" if toc else ""
    if len(md) <= MAX_TOOL_CHARS:
        return head + md
    if not find:
        return _clip(head + toc_text + md)
    chunks = [("", s) for s in sections]  # same texts as indexed in the background
    picked = _best_windows(chunks, find, k=8)
    body, used = [], 0
    for _, text in picked:
        text = _clip(text, 8000)
        if used + len(text) > MAX_TOOL_CHARS:
            break
        body.append(text)
        used += len(text)
    if not body:
        return _clip(head + toc_text + md)
    return head + toc_text + f"Sections about '{find}':\n\n" + "\n\n".join(body)


def web_search(query: str) -> str:
    data = _firecrawl("search", {"query": query, "limit": 5})
    rows = data.get("web", data if isinstance(data, list) else [])
    out = [f"- {r.get('title', '')}\n  {r.get('url', '')}\n  {r.get('description', '')}" for r in rows]
    return "\n".join(out) or "No results."


# ---- sources ---------------------------------------------------------------------

class Source:
    kind = "screen"
    label = ""

    def __init__(self):
        self._loaded: Future = _pool.submit(self._load)

    def _load(self) -> None:
        pass

    def wait(self, timeout: float | None = None) -> bool:
        try:
            self._loaded.result(timeout=timeout)
            return True
        except Exception:
            return False

    def load_error(self) -> str | None:
        if self._loaded.done() and self._loaded.exception():
            return str(self._loaded.exception())
        return None

    @property
    def loading(self) -> bool:
        return not self._loaded.done()

    def digest(self, screen_text: str) -> str:
        return ""

    def tools(self) -> list[type[BaseModel]]:
        return []

    def call(self, args: BaseModel) -> str:
        raise NotImplementedError


class YouTubeSource(Source):
    kind = "youtube"

    def __init__(self, video_id: str, url: str, window_title: str):
        self.video_id, self.url = video_id, url
        self.label = re.sub(r"\s*-\s*YouTube.*$", "", window_title).strip() or "YouTube video"
        self.segments: list[tuple[float, str]] = []
        self.description = self.channel = ""
        self.length = 0
        self.language = ""
        self.start_hint = self._t_param(url)
        super().__init__()

    @staticmethod
    def _t_param(url: str) -> float | None:
        t = urllib.parse.parse_qs(urllib.parse.urlparse(url).query).get("t", [None])[0]
        if not t:
            return None
        m = re.fullmatch(r"(?:(\d+)h)?(?:(\d+)m)?(?:(\d+)s?)?", t)
        return float(int(m[1] or 0) * 3600 + int(m[2] or 0) * 60 + int(m[3] or 0)) if m else None

    def _load(self) -> None:
        details = threading.Thread(target=self._load_details_quietly, daemon=True)
        details.start()
        try:
            self._load_captions()
        finally:
            details.join(15)
        index(self._chunks())

    def _load_details_quietly(self) -> None:
        try:
            self._load_details()
        except Exception:
            pass  # title/description are nice to have; captions are what matter

    def _load_captions(self) -> None:
        from youtube_transcript_api import YouTubeTranscriptApi

        api = YouTubeTranscriptApi()
        try:
            fetched = api.fetch(self.video_id, languages=["en", "en-US", "en-GB", "bn", "hi"])
        except Exception:
            fetched = next(iter(api.list(self.video_id))).fetch()  # any language there is
        self.language = fetched.language_code
        self.segments = [(s.start, s.text.replace("\n", " ")) for s in fetched.snippets]

    def _load_details(self) -> None:
        r = httpx.get(f"https://www.youtube.com/watch?v={self.video_id}", timeout=15, follow_redirects=True,
                      headers={"Accept-Language": "en-US,en;q=0.9", "User-Agent": "Mozilla/5.0"})
        html = r.text
        m = re.search(r'"shortDescription":"((?:\\.|[^"\\])*)"', html)
        if m:
            self.description = json.loads(f'"{m.group(1)}"')
        m = re.search(r'"lengthSeconds":"(\d+)"', html)
        self.length = int(m.group(1)) if m else 0
        m = re.search(r'"author":"((?:\\.|[^"\\])*)"', html)
        self.channel = json.loads(f'"{m.group(1)}"') if m else ""

    def transcript(self, start: float = 0, end: float | None = None) -> str:
        """Captions grouped into ~20 s lines: '[3:05] ...'."""
        out, cur, cur_t = [], [], None
        for t, text in self.segments:
            if t < start or (end is not None and t > end):
                continue
            if cur_t is None:
                cur_t = t
            cur.append(text)
            if t - cur_t >= 20:
                out.append(f"[{_mmss(cur_t)}] {' '.join(cur)}")
                cur, cur_t = [], None
        if cur:
            out.append(f"[{_mmss(cur_t)}] {' '.join(cur)}")
        return "\n".join(out)

    def _chunks(self) -> list[tuple[str, str]]:
        """~30 s windows of captions, labelled with their start time."""
        chunks, cur, t0 = [], [], None
        for t, text in self.segments:
            t0 = t if t0 is None else t0
            cur.append(text)
            if t - t0 >= 30:
                chunks.append((_mmss(t0), " ".join(cur)))
                cur, t0 = [], None
        if cur:
            chunks.append((_mmss(t0), " ".join(cur)))
        return chunks

    def current_time(self, screen_text: str) -> float | None:
        m = re.search(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\s*/\s*(\d{1,2}:\d{2}(?::\d{2})?)\b", screen_text)
        if m:
            return _seconds(m.group(1))
        return self.start_hint

    def digest(self, screen_text: str) -> str:
        lines = [f"YouTube video: {self.label}" + (f" by {self.channel}" if self.channel else "")]
        if self.length:
            lines.append(f"Length {_mmss(self.length)}.")
        if self.description:
            lines.append("Description: " + _clip(self.description, 1200))
        if self.loading:
            lines.append("Captions: still downloading; the transcript tools will wait for them.")
            return "\n".join(lines)
        if not self.segments:
            lines.append("Captions: not available" + (f" ({self.load_error()})" if self.load_error() else "") + ".")
            return "\n".join(lines)
        lines.append(f"Captions: available ({self.language}); full transcript via tools.")
        now = self.current_time(screen_text)
        if now is not None:
            lines.append(f"The player is at about {_mmss(now)}. Captions from the last 2.5 minutes:")
            lines.append(self.transcript(max(0, now - 150), now + 10) or "(none)")
        else:
            lines.append("Current playback position unknown (player controls hidden): search the transcript for what's on screen.")
        return "\n".join(lines)

    def tools(self):
        return [GetVideoTranscript, SearchVideoTranscript] if self.segments or self.loading else []

    def call(self, args):
        self.wait()
        if isinstance(args, GetVideoTranscript):
            return _clip(self.transcript(args.start_sec, args.end_sec) or "No captions in that range.")
        if isinstance(args, SearchVideoTranscript):
            hits = _best_windows(self._chunks(), args.query)
            return "\n".join(f"[{ts}] {text}" for ts, text in hits) or "No part of the video mentions that."
        raise ValueError("not a video tool")


class DocumentSource(Source):
    kind = "document"

    def __init__(self, path: str):
        self.path = path
        self.label = os.path.basename(path)
        self.pages: list[str] = []
        self.unit = "slide" if path.lower().endswith(".pptx") else "page"
        super().__init__()

    def _load(self) -> None:
        p = self.path.lower()
        if p.endswith(".pdf"):
            from pypdf import PdfReader

            self.pages = [(pg.extract_text() or "").strip() for pg in PdfReader(self.path).pages]
        elif p.endswith(".pptx"):
            from pptx import Presentation

            for slide in Presentation(self.path).slides:
                parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
                self.pages.append("\n".join(parts))
        else:
            if p.endswith(".docx"):
                from docx import Document

                text = "\n".join(par.text for par in Document(self.path).paragraphs)
            else:
                text = Path(self.path).read_text(encoding="utf-8", errors="replace")
            self.pages = [text[i : i + 3000] for i in range(0, len(text), 3000)] or [""]
        index(self._chunks())

    def _chunks(self) -> list[tuple[str, str]]:
        """~900-character overlapping passages, labelled with their page. Searching
        passages (not whole pages) returns the exact part that answers the question."""
        out = []
        for i, pg in enumerate(self.pages):
            for start in range(0, max(1, len(pg)), 750):
                part = pg[start:start + 900].strip()
                if part:
                    out.append((f"{self.unit} {i + 1}", part))
        return out

    def current_page(self, screen_text: str) -> int | None:
        s = _words(screen_text)
        if not s or not self.pages:
            return None
        scores = [len(s & _words(pg)) / (len(s) + 1) for pg in self.pages]
        best = max(range(len(scores)), key=scores.__getitem__)
        return best + 1 if scores[best] >= 0.12 else None

    def _pages(self, a: int, b: int) -> str:
        a, b = max(1, a), min(len(self.pages), b)
        return "\n\n".join(f"--- {self.unit} {i} ---\n{self.pages[i - 1]}" for i in range(a, b + 1))

    def digest(self, screen_text: str) -> str:
        if self.loading:
            return f"Open document: {self.label} (still loading; the document tools will wait for it)."
        if not self.pages:
            err = self.load_error()
            return f"Open document: {self.label} (could not read it{': ' + err if err else ''})."
        lines = [f"Open document: {self.label}, {len(self.pages)} {self.unit}s. Full text via tools."]
        cur = self.current_page(screen_text)
        if cur:
            lines.append(f"The screen shows {self.unit} {cur}. Text of {self.unit}s {max(1, cur - 1)}-{min(len(self.pages), cur + 1)}:")
            lines.append(_clip(self._pages(cur - 1, cur + 1), 7000))
        else:
            lines.append(f"Couldn't tell which {self.unit} is on screen.")
        return "\n".join(lines)

    def tools(self):
        return [ReadDocumentPages, SearchDocument] if self.pages or self.loading else []

    def call(self, args):
        self.wait()
        if isinstance(args, ReadDocumentPages):
            a, b = sorted((args.start_page, args.end_page))
            return _clip(self._pages(a, min(b, a + 19)))
        if isinstance(args, SearchDocument):
            hits = _best_windows(self._chunks(), args.query, k=8)
            return "\n\n".join(f"--- {lab} ---\n{txt}" for lab, txt in hits) or "Not found in the document."
        raise ValueError("not a document tool")


class WebPageSource(Source):
    kind = "webpage"

    def __init__(self, url: str, window_title: str):
        self.url = url
        self.label = re.sub(r"\s*-\s*(Google Chrome|Microsoft​? Edge|Mozilla Firefox|Brave).*$", "", window_title)
        super().__init__()

    def _load(self) -> None:
        # Read (and index) the whole page while the user types, so a lookup is instant.
        if config.PREFETCH_PAGES and config.FIRECRAWL_KEY:
            _, md = _scrape(self.url)
            index([("", s) for s in _sections(md)])

    def digest(self, screen_text: str) -> str:
        return f"Web page: {self.label}\n{self.url}\n(Only the visible part is on screen; the full page is available via read_webpage.)"


# ---- detection -------------------------------------------------------------------

def _youtube_id(url: str) -> str | None:
    m = re.search(r"(?:youtube\.com/(?:watch\?(?:.*&)?v=|shorts/|live/)|youtu\.be/)([\w-]{11})", url)
    return m.group(1) if m else None


def _resolve_lnk(lnk: Path) -> str | None:
    cmd = f"(New-Object -ComObject WScript.Shell).CreateShortcut('{lnk}').TargetPath"
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-Command", cmd], capture_output=True, text=True,
                             timeout=5, creationflags=0x08000000).stdout.strip()
    except Exception:
        return None
    return out if out and os.path.isfile(out) else None


def locate_file(name: str, exts: tuple[str, ...]) -> str | None:
    """Find a file on disk from the name in a window title."""
    candidates = [name] if name.lower().endswith(DOC_EXTS) else [name + e for e in exts]
    recent = Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Windows" / "Recent"
    for c in candidates:
        for lnk in (recent / f"{c}.lnk", recent / f"{Path(c).stem}.lnk"):
            if lnk.exists():
                target = _resolve_lnk(lnk)
                if target and target.lower().endswith(exts):
                    return target
    home = Path.home()
    roots = [home / "Downloads", home / "Desktop", home / "Documents", home / "OneDrive", Path.cwd().parent]
    wanted = {c.lower() for c in candidates}
    for root in roots:
        if not root.exists():
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            if dirpath.count(os.sep) - str(root).count(os.sep) >= 4:
                dirnames[:] = []
            for f in filenames:
                if f.lower() in wanted:
                    return os.path.join(dirpath, f)
    return None


APP_DOCS = {  # desktop apps whose title is "<file name> - <App>"
    "powerpnt.exe": (".pptx",), "winword.exe": (".docx",), "acrord32.exe": (".pdf",), "acrobat.exe": (".pdf",),
    "sumatrapdf.exe": (".pdf",), "foxitpdfreader.exe": (".pdf",), "foxitreader.exe": (".pdf",),
    "notepad.exe": (".txt", ".md"), "code.exe": (".md", ".txt", ".pdf"),
}


def _download(url: str) -> str:
    fd, path = tempfile.mkstemp(suffix=".pdf", prefix="chalktalk-")
    with httpx.stream("GET", url, timeout=30, follow_redirects=True) as r, os.fdopen(fd, "wb") as f:
        r.raise_for_status()
        for chunk in r.iter_bytes():
            f.write(chunk)
    return path


_cache: dict[str, Source] = {}  # asking again about the same video/file reuses what's loaded


def _cached(key: str, make: Callable[[], Source]) -> Source:
    src = _cache.get(key)
    if src is None or src.load_error():
        src = _cache[key] = make()
        while len(_cache) > 8:
            _cache.pop(next(iter(_cache)))
    return src


def detect(fg: Foreground) -> Source | None:
    """Work out what the user is looking at and start loading it."""
    url = browser_url(fg.hwnd) if fg.is_browser else None
    if url:
        vid = _youtube_id(url)
        if vid:
            src = _cached(f"yt:{vid}", lambda: YouTubeSource(vid, url, fg.title))
            src.start_hint = YouTubeSource._t_param(url)
            return src
        path = urllib.parse.urlparse(url).path
        if url.startswith("file:///") and path.lower().endswith(DOC_EXTS):
            local = urllib.parse.unquote(url[len("file:///"):])
            return _cached(f"doc:{local}:{os.path.getmtime(local)}", lambda: DocumentSource(local))
        if not url.startswith(("http://", "https://")):
            return None  # a local image or file that isn't a document: the screen is the source
        if path.lower().endswith(".pdf"):
            return _cached(f"url:{url}", lambda: DocumentSource(_download(url)))
        return WebPageSource(url, fg.title)
    exts = APP_DOCS.get(fg.exe)
    if exts:
        name = re.split(r"\s+[-–—]\s+", fg.title)[0].strip("*● ").strip()
        path = locate_file(name, exts) if name else None
        if path:
            return _cached(f"doc:{path}:{os.path.getmtime(path)}", lambda: DocumentSource(path))
    return None


# ---- the toolbox the model sees -----------------------------------------------

class Toolbox:
    def __init__(self, source: Source | None):
        self.source = source
        self.models: dict[str, type[BaseModel]] = {}
        for m in (source.tools() if source else []):
            self.models[_name(m)] = m
        if isinstance(source, WebPageSource):
            self.models[_name(ReadWebpage)] = ReadWebpage
        if config.FIRECRAWL_KEY:
            self.models[_name(WebSearch)] = WebSearch
            self.models.setdefault(_name(ReadWebpage), ReadWebpage)

    def specs(self) -> list[dict]:
        return [openai.pydantic_function_tool(m, name=n, description=(m.__doc__ or "").strip())
                for n, m in self.models.items()]

    def parse(self, name: str, arguments: str) -> BaseModel:
        return self.models[name].model_validate_json(arguments or "{}")

    def run(self, args: BaseModel) -> str:
        try:
            if isinstance(args, WebSearch):
                return web_search(args.query)
            if isinstance(args, ReadWebpage):
                url = args.url if args.url and args.url.lower() not in ("null", "none") else None
                url = url or (self.source.url if isinstance(self.source, WebPageSource) else None)
                return read_url(url, args.find) if url else "No page given."
            return self.source.call(args)
        except (httpx.HTTPError, RuntimeError, ValueError, OSError) as e:
            return f"Tool failed: {e}"


def _name(model: type[BaseModel]) -> str:
    return re.sub(r"(?<!^)(?=[A-Z])", "_", model.__name__).lower()  # GetVideoTranscript -> get_video_transcript
